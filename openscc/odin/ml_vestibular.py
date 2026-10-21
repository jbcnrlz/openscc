import csv
import io
import json
import math
from datetime import timedelta
from pathlib import Path
import numpy as np
from scipy.interpolate import interp1d

BASE_CSV_PATH = Path(__file__).resolve().parent / "data" / "vestibular.csv"


class PreditorVestibularOdin:
    """
    Motor de Aprendizado de Máquina para Projeção de Demanda do Vestibular no Odin.
    Combina:
      1. Aprendizado de Curvas Históricas por Turno/Curso (Prior Histórico Ponderado)
      2. Regressão Funcional por Similaridade de Trajetória (Kernel Trajectory Matching / k-NN Funcional)
         usando o "andamento" diário real da campanha (CampaignDailyRecord)
      3. Regressão Supervisionada Regularizada (Ridge Trajectory Regressor) sobre features
         dinâmicas do andamento (velocidade recente, aceleração, convexidade AUC e turno)
      4. Integração com o Funil de Leads (CRM) e Controle Ótimo de Patrocínios (Yield Management)
    """

    GRID_SIZE = 101  # Grade normalizada tau in [0.00, 0.01, ..., 1.00]

    def __init__(self, caminho_csv=None, pesos_edicoes=None):
        self.caminho_csv = Path(caminho_csv) if caminho_csv else BASE_CSV_PATH
        self.pesos = pesos_edicoes if pesos_edicoes is not None else [1.0, 2.0, 3.0]
        self.grid = np.linspace(0.0, 1.0, self.GRID_SIZE)
        self.series_treino = []
        self.series_parciais = []
        self.dados_treino = self._treinar_curvas()

    @staticmethod
    def detectar_turno_e_curso(nome_curso: str = "", nome_campanha: str = "", periodo_override: str = None):
        """
        Infere automaticamente o turno ('M', 'N' ou 'G') e a sigla histórica ('AI', 'TI', 'GE' ou 'NOVO')
        a partir do nome do curso e/ou da campanha, respeitando eventual override.
        """
        texto = f"{nome_curso or ''} {nome_campanha or ''}".lower()

        # Detecção de sigla histórica
        sigla = "NOVO"
        if any(k in texto for k in ["automação", "automacao", "ai ", " ai", "(ai)"]):
            sigla = "AI"
        elif any(k in texto for k in ["tecnologia da informação", "tecnologia da informacao", "gti", "ti ", " ti", "(ti)"]):
            sigla = "TI"
        elif any(k in texto for k in ["gestão empresarial", "gestao empresarial", "ge ", " ge", "(ge)"]):
            sigla = "GE"

        if periodo_override and periodo_override.upper() in ["M", "N", "G"]:
            periodo = periodo_override.upper()
        else:
            if any(k in texto for k in ["matutino", "manhã", "manha", "(m)", "- m", "dsm", "desenvolvimento de software", "engenharia de alimentos"]):
                periodo = "M"
            elif any(k in texto for k in ["noturno", "noite", "(n)", "- n", "tecnologia de alimentos", "gestão comercial", "gestao comercial", "automação", "automacao", "tecnologia da informação"]):
                periodo = "N"
            else:
                periodo = "N" if sigla in ["AI", "TI"] else "G"

        turno_extenso = (
            "Matutino" if periodo == "M" else ("Noturno" if periodo == "N" else "Geral (Média Unidade)")
        )
        return periodo, sigla, turno_extenso

    def _treinar_curvas(self):
        """
        Lê a base histórica de treino (vestibular.csv) e constrói tanto a estrutura compatível
        com o PreditorVestibularGenerico quanto as curvas funcionais normalizadas para o Ensemble ML.
        """
        self.series_treino = []
        self.series_parciais = []

        if not self.caminho_csv.exists():
            return []

        with open(self.caminho_csv, mode="r", encoding="utf-8") as f:
            reader = list(csv.reader(f))

        if not reader:
            return []

        header = [h.strip() for h in reader[0]]
        col_dias_indices = [i for i, h in enumerate(header) if h not in ["curso", "período", "periodo"]]
        max_col_dia = len(col_dias_indices) - 1

        raw_rows = []
        for row_idx, r in enumerate(reader[1:]):
            if not r or len(r) < 3:
                continue
            curso = r[0].strip()
            periodo = r[1].strip().upper()

            # Descobre até qual coluna a linha possui dados reais preenchidos
            last_valid_day = 0
            raw_vals = []
            for d_idx, col_i in enumerate(col_dias_indices):
                val_str = r[col_i].strip() if col_i < len(r) else ""
                if val_str not in ("", "?"):
                    try:
                        v_num = float(val_str)
                        last_valid_day = d_idx
                        raw_vals.append(v_num)
                    except ValueError:
                        raw_vals.append(np.nan)
                else:
                    raw_vals.append(np.nan)

            # Forward-fill e zero-fill inicial
            filled_vals = []
            curr = 0.0
            for v in raw_vals:
                if not np.isnan(v):
                    curr = v
                filled_vals.append(curr)

            raw_rows.append({
                "row_idx": row_idx,
                "curso": curso,
                "período": periodo,
                "last_valid_day": last_valid_day,
                "vals": np.array(filled_vals, dtype=float),
            })

        # As 3 edições completas padrão (linhas 0..3 -> 63 dias, 4..7 -> 44 dias, 8..11 -> 44 dias)
        # Se novas linhas completas forem adicionadas ao CSV, também são incorporadas automaticamente.
        edicoes_spec = [
            {"slice": slice(0, 4), "dias_totais": 63, "peso": self.pesos[0] if len(self.pesos) > 0 else 1.0, "edicao": 1},
            {"slice": slice(4, 8), "dias_totais": 44, "peso": self.pesos[1] if len(self.pesos) > 1 else 2.0, "edicao": 2},
            {"slice": slice(8, 12), "dias_totais": 44, "peso": self.pesos[2] if len(self.pesos) > 2 else 3.0, "edicao": 3},
        ]

        edicoes_treino = []
        for spec in edicoes_spec:
            sub_rows = raw_rows[spec["slice"]]
            if not sub_rows:
                continue
            t_max = min(spec["dias_totais"], max_col_dia)
            edicoes_treino.append({
                "rows": sub_rows,
                "dias_totais": t_max,
                "peso": spec["peso"],
                "edicao": spec["edicao"],
            })

            for item in sub_rows:
                vals_ed = item["vals"][: t_max + 1]
                final_val = float(vals_ed[-1])
                if final_val > 0:
                    t_norm = np.linspace(0.0, 1.0, t_max + 1)
                    f_interp = interp1d(t_norm, vals_ed, kind="linear", bounds_error=False, fill_value=(vals_ed[0], final_val))
                    raw_grid = f_interp(self.grid)
                    curve_grid = np.clip(raw_grid / final_val, 0.0, 1.0)
                    self.series_treino.append({
                        "id": f"{item['curso']}-{item['período']}-Ed{spec['edicao']}",
                        "label": f"{item['curso']} ({item['período']}) - {spec['edicao']}ª Ed.",
                        "curso": item["curso"],
                        "periodo": item["período"],
                        "edicao": spec["edicao"],
                        "dias_totais": t_max,
                        "peso": float(spec["peso"]),
                        "vals": vals_ed,
                        "final": final_val,
                        "raw_grid": raw_grid,
                        "curve": curve_grid,
                    })

        # Linhas adicionais no CSV (ex: linhas 12..15 que estão em andamento até o dia 33 ou novas edições completas)
        if len(raw_rows) > 12:
            for item in raw_rows[12:]:
                lvd = item["last_valid_day"]
                vals_sub = item["vals"][: lvd + 1]
                if lvd >= 40 and vals_sub[-1] > 0:
                    # Trata como uma edição completa extra se tiver >= 40 dias
                    t_norm = np.linspace(0.0, 1.0, lvd + 1)
                    f_interp = interp1d(t_norm, vals_sub, kind="linear", bounds_error=False, fill_value=(vals_sub[0], vals_sub[-1]))
                    raw_grid = f_interp(self.grid)
                    self.series_treino.append({
                        "id": f"{item['curso']}-{item['período']}-Extra",
                        "label": f"{item['curso']} ({item['período']}) - Ed. Extra",
                        "curso": item["curso"],
                        "periodo": item["período"],
                        "edicao": 4,
                        "dias_totais": lvd,
                        "peso": 3.5,
                        "vals": vals_sub,
                        "final": float(vals_sub[-1]),
                        "raw_grid": raw_grid,
                        "curve": np.clip(raw_grid / float(vals_sub[-1]), 0.0, 1.0),
                    })
                elif lvd >= 10:
                    self.series_parciais.append({
                        "curso": item["curso"],
                        "periodo": item["período"],
                        "dias_observados": lvd,
                        "vals": vals_sub,
                    })

        return edicoes_treino

    def obter_ritmo_historico(self, periodo: str, dia_atual: int, total_dias: int):
        """
        Extrai o ritmo de conversão aprendido para o turno correspondente ('M' ou 'N').
        Se o período for diferente, utiliza a curva média ponderada da unidade.
        Mantém total compatibilidade com vestibularEstimator.py.
        """
        total_dias = max(1, int(total_dias))
        dia_atual = max(0, min(int(dia_atual), total_dias))
        frac = min(dia_atual / total_dias, 1.0)
        props, pesos_aplicados = [], []

        for ed in self.dados_treino:
            rows = ed["rows"]
            if periodo in ["M", "N"]:
                filtradas = [r for r in rows if r["período"] == periodo]
                rows_usar = filtradas if filtradas else rows
            else:
                rows_usar = rows

            t_max = ed["dias_totais"]
            dia_ref = min(int(round(frac * t_max)), t_max)

            for r in rows_usar:
                fin = float(r["vals"][t_max])
                val = float(r["vals"][dia_ref])
                if fin > 0:
                    props.append(val / fin)
                    pesos_aplicados.append(ed["peso"])

        if not props:
            return max(frac, 0.05), max(frac * 0.8, 0.05), min(frac * 1.2, 1.0)

        props = np.array(props, dtype=float)
        pesos_aplicados = np.array(pesos_aplicados, dtype=float)

        p_ponderado = float(np.average(props, weights=pesos_aplicados))
        p_min = float(np.min(props))  # Menor percentual histórico => Cenário Otimista
        p_max = float(np.max(props))  # Maior percentual histórico => Cenário Pessimista
        return p_ponderado, p_min, p_max

    @staticmethod
    def _extrair_features_trajetoria(y_obs_grid: np.ndarray, periodo: str):
        """
        Extrai vetor de atributos (features) da trajetória acumulada observada até o momento:
          1. Intercepto (1.0)
          2. Indicador de Turno Matutino (is_M)
          3. Área sob a curva normalizada (AUC / Convexidade do andamento)
          4. Velocidade recente (últimos 25% do tempo decorrido)
          5. Velocidade de médio prazo (últimos 50% do tempo decorrido)
        """
        y_now = max(float(y_obs_grid[-1]), 1e-3)
        n = len(y_obs_grid)
        y_norm = y_obs_grid / y_now
        auc = float(np.mean(y_norm))
        k25 = max(1, int(n * 0.25))
        k50 = max(1, int(n * 0.50))
        vel25 = float((y_norm[-1] - y_norm[-1 - k25]) / max(k25 / n, 1e-3))
        vel50 = float((y_norm[-1] - y_norm[-1 - k50]) / max(k50 / n, 1e-3))
        is_m = 1.0 if periodo == "M" else 0.0
        return np.array([1.0, is_m, auc, vel25, vel50], dtype=float)

    def _projetar_ensemble_andamento(
        self,
        periodo: str,
        sigla_curso: str,
        inscritos_hoje: float,
        dia_atual: int,
        total_dias: int,
        andamento_diario=None,
        num_registros_reais: int = 0,
    ):
        """
        Executa os modelos de Aprendizado de Máquina utilizando tanto o ponto atual quanto
        toda a série temporal de andamento (quando disponível no Odin).
        """
        total_dias = max(1, int(total_dias))
        dia_atual = max(0, min(int(dia_atual), total_dias))
        frac = min(dia_atual / total_dias, 1.0)
        tau_idx = min(int(round(frac * (self.GRID_SIZE - 1))), self.GRID_SIZE - 1)

        # 1. Estimador Pontual Base (vestibularEstimator.py)
        p_med, p_min, p_max = self.obter_ritmo_historico(periodo, dia_atual, total_dias)
        proj_pontual = inscritos_hoje / p_med if p_med > 0.01 else float(inscritos_hoje)
        proj_otimista_hist = inscritos_hoje / p_min if p_min > 0.01 else float(inscritos_hoje)
        proj_pessimista_hist = inscritos_hoje / p_max if p_max > 0.01 else float(inscritos_hoje)

        # Curva histórica média ponderada do turno (para fallback ou prior)
        pesos_prior = []
        curvas_prior = []
        for s in self.series_treino:
            w_p = s["peso"] * (2.5 if s["periodo"] == periodo else (1.0 if periodo not in ["M", "N"] else 0.4))
            pesos_prior.append(w_p)
            curvas_prior.append(s["curve"])
        curva_prior_media = np.average(curvas_prior, axis=0, weights=pesos_prior) if curvas_prior else np.linspace(0, 1, self.GRID_SIZE)

        # Verifica se temos andamento diário válido (série com >= 2 dias e inscritos > 0)
        tem_andamento = (
            andamento_diario is not None
            and len(andamento_diario) >= 2
            and tau_idx >= 3
            and inscritos_hoje > 0
            and (num_registros_reais >= 2 or len(set(andamento_diario)) >= 2)
        )

        if not tem_andamento or not self.series_treino:
            # Sem histórico diário detalhado: usa o modelo calibrado por turno histórico
            confianca = round(min(95.0, 25.0 + 55.0 * frac), 1) if dia_atual > 0 else 0.0
            return {
                "usou_andamento_ml": False,
                "projecao_esperada": max(float(inscritos_hoje), proj_pontual),
                "projecao_pontual": max(float(inscritos_hoje), proj_pontual),
                "projecao_kernel": max(float(inscritos_hoje), proj_pontual),
                "projecao_ridge": max(float(inscritos_hoje), proj_pontual),
                "cenario_pessimista": max(float(inscritos_hoje), proj_pessimista_hist),
                "cenario_otimista": max(float(inscritos_hoje), proj_otimista_hist),
                "ritmo_medio": p_med,
                "ritmo_ajustado_ml": p_med,
                "curva_aprendida_grid": curva_prior_media,
                "curva_mais_similar": f"Curva Média Histórica ({'Matutino' if periodo == 'M' else ('Noturno' if periodo == 'N' else 'Geral')})",
                "similaridade_max_pct": round(min(98.0, 70.0 + 25.0 * frac), 1),
                "top_curvas_similares": [],
                "aceleracao_ratio": 1.0,
                "aceleracao_label": "Dentro do Padrão Histórico",
                "confianca_modelo": confianca,
            }

        # Interpola o andamento observado do dia 0..dia_atual para a grade normalizada 0..tau_idx
        y_arr = np.array(andamento_diario, dtype=float)
        t_obs_norm = np.linspace(0.0, self.grid[tau_idx], len(y_arr))
        f_obs = interp1d(t_obs_norm, y_arr, kind="linear", bounds_error=False, fill_value=(y_arr[0], y_arr[-1]))
        y_obs_grid = f_obs(self.grid[: tau_idx + 1])
        y_now = max(float(y_obs_grid[-1]), 1.0)
        y_shape = y_obs_grid / y_now

        # Pesos temporais exponenciais (valoriza o andamento mais recente)
        tw = np.exp(np.linspace(-2.0, 0.0, tau_idx + 1))

        # =====================================================================
        # MODELO 1: Kernel Trajectory Matching + Escala de Velocidade Recente
        # =====================================================================
        preds_kernel = []
        weights_kernel = []
        sim_scores = []
        k_back = max(2, int(tau_idx * 0.30))
        dy_rec = max(0.0, float(y_obs_grid[-1] - y_obs_grid[-1 - k_back]))
        slope_y = float(dy_rec / y_now)

        for s in self.series_treino:
            c_obs = s["curve"][: tau_idx + 1]
            c_end = max(float(c_obs[-1]), 0.03)
            c_shape = c_obs / c_end

            # Distância funcional ponderada da forma da curva (L2 com decaimento exponencial)
            dist_shape = float(np.sqrt(np.average((y_shape - c_shape) ** 2, weights=tw)))
            # Diferença de velocidade/inclinação recente
            dc_rec = max(0.02, float(c_obs[-1] - c_obs[-1 - k_back]))
            slope_c = float(dc_rec / c_end)
            diff_slope = abs(slope_y - slope_c)

            # Projeção combinando escala de nível atual e escala de velocidade recente do andamento
            f_level = float(y_now / c_end)
            f_vel = float(dy_rec / dc_rec)
            f_scale = 0.55 * f_level + 0.45 * float(np.clip(f_vel, 0.25 * f_level, 1.85 * f_level))
            f_est = max(y_now, y_now + f_scale * (1.0 - c_end))

            # Afinidade por turno, curso e kernel radial de trajetória
            w_shift = 2.5 if s["periodo"] == periodo else (1.0 if periodo not in ["M", "N"] else 0.5)
            w_curso = 1.5 if s["curso"] == sigla_curso else 1.0
            sim_kernel = float(np.exp(-5.0 * dist_shape - 4.0 * diff_slope))
            w_total = float(s["peso"] * w_shift * w_curso * sim_kernel)

            preds_kernel.append(f_est)
            weights_kernel.append(w_total)
            sim_pct = round(100.0 * float(np.exp(-2.5 * dist_shape - 1.2 * diff_slope)), 1)
            sim_scores.append({
                "label": s["label"],
                "similaridade_pct": sim_pct,
                "projecao_curva": int(round(f_est)),
                "peso_final": w_total,
                "curve": s["curve"],
            })

        preds_kernel = np.array(preds_kernel, dtype=float)
        weights_kernel = np.array(weights_kernel, dtype=float)
        proj_kernel = float(np.average(preds_kernel, weights=weights_kernel))

        # Curva funcional aprendida (combinação convexa das curvas históricas ponderada pela similaridade do andamento)
        curva_aprendida_grid = np.average(
            [s["curve"] for s in self.series_treino], axis=0, weights=weights_kernel
        )

        # Ordena as curvas históricas mais similares ao andamento atual
        sim_scores_sorted = sorted(sim_scores, key=lambda x: x["peso_final"], reverse=True)
        melhor_curva = sim_scores_sorted[0] if sim_scores_sorted else None

        # Cálculo da razão de aceleração recente vs esperado histórico na mesma janela
        slope_hist = float(
            (curva_prior_media[tau_idx] - curva_prior_media[max(0, tau_idx - k_back)])
            / max(curva_prior_media[tau_idx], 0.02)
        )
        aceleracao_ratio = round(slope_y / max(slope_hist, 0.02), 2)

        # =====================================================================
        # MODELO 2: Ridge Regression Supervisionada + Ajuste de Momento
        # =====================================================================
        X_tr, y_tr, w_tr = [], [], []
        for s in self.series_treino:
            s_obs = s["raw_grid"][: tau_idx + 1]
            if s_obs[-1] > 0:
                X_tr.append(self._extrair_features_trajetoria(s_obs, s["periodo"]))
                y_tr.append(math.log(max(s["final"] / s_obs[-1], 1.0)))
                w_shift = 2.0 if s["periodo"] == periodo else (1.0 if periodo not in ["M", "N"] else 0.7)
                w_tr.append(s["peso"] * w_shift)

        if len(X_tr) >= 4:
            X_tr_arr = np.array(X_tr, dtype=float)
            y_tr_arr = np.array(y_tr, dtype=float)
            W_mat = np.diag(w_tr)
            reg_mat = np.diag([0.0, 1.5, 4.0, 4.0, 4.0])
            try:
                beta = np.linalg.solve(X_tr_arr.T @ W_mat @ X_tr_arr + reg_mat, X_tr_arr.T @ W_mat @ y_tr_arr)
                x_te = self._extrair_features_trajetoria(y_obs_grid, periodo)
                log_mult = float(np.clip(x_te @ beta, 0.0, 3.5))
                proj_ridge_base = float(max(y_now, y_now * math.exp(log_mult)))
                # Modula o ganho futuro pela aceleração observada do andamento atual
                fator_mom = float(np.clip(0.75 + 0.25 * aceleracao_ratio, 0.65, 1.45))
                proj_ridge = y_now + (proj_ridge_base - y_now) * fator_mom
            except np.linalg.LinAlgError:
                proj_ridge = proj_kernel
        else:
            proj_ridge = proj_kernel

        # =====================================================================
        # COMBINAÇÃO ENSEMBLE (Ponderação Adaptativa pelo Volume de Andamento)
        # =====================================================================
        peso_andamento = min(1.0, max(0.35, (num_registros_reais if num_registros_reais > 0 else len(y_arr)) / 6.0))
        proj_ml_puro = 0.70 * proj_kernel + 0.30 * proj_ridge
        proj_final = peso_andamento * proj_ml_puro + (1.0 - peso_andamento) * proj_pontual
        proj_final = max(float(inscritos_hoje), proj_final)

        if aceleracao_ratio >= 1.15:
            aceleracao_label = "Em Aceleração (Acima do Ritmo)"
        elif aceleracao_ratio <= 0.85:
            aceleracao_label = "Desacelerando (Atenção ao Ritmo)"
        else:
            aceleracao_label = "Ritmo Consistente com Histórico"

        # Intervalo de Confiança / Cenários calibrados pelo Ensemble + Quantis Ponderados
        desvio_ponderado = float(np.sqrt(np.average((preds_kernel - proj_kernel) ** 2, weights=weights_kernel)))
        cenario_pessimista = max(
            float(inscritos_hoje),
            0.5 * proj_pessimista_hist + 0.5 * (proj_final - 1.15 * desvio_ponderado),
        )
        cenario_otimista = max(
            proj_final,
            0.5 * proj_otimista_hist + 0.5 * (proj_final + 1.25 * desvio_ponderado),
        )

        ritmo_ajustado = float(inscritos_hoje / proj_final) if proj_final > 0 else p_med
        confianca = round(
            min(98.0, 35.0 + 45.0 * frac + min(18.0, (num_registros_reais or len(y_arr)) * 1.5)),
            1,
        )

        return {
            "usou_andamento_ml": True,
            "projecao_esperada": proj_final,
            "projecao_pontual": max(float(inscritos_hoje), proj_pontual),
            "projecao_kernel": max(float(inscritos_hoje), proj_kernel),
            "projecao_ridge": max(float(inscritos_hoje), proj_ridge),
            "cenario_pessimista": cenario_pessimista,
            "cenario_otimista": cenario_otimista,
            "ritmo_medio": p_med,
            "ritmo_ajustado_ml": ritmo_ajustado,
            "curva_aprendida_grid": curva_aprendida_grid,
            "curva_mais_similar": melhor_curva["label"] if melhor_curva else "Média Histórica",
            "similaridade_max_pct": melhor_curva["similaridade_pct"] if melhor_curva else 80.0,
            "top_curvas_similares": [
                {
                    "label": item["label"],
                    "similaridade_pct": item["similaridade_pct"],
                    "projecao": item["projecao_curva"],
                }
                for item in sim_scores_sorted[:3]
            ],
            "aceleracao_ratio": aceleracao_ratio,
            "aceleracao_label": aceleracao_label,
            "confianca_modelo": confianca,
        }

    def projetar(
        self,
        nome_curso: str,
        periodo: str = None,
        inscritos_hoje: float = 0,
        dia_atual: int = 1,
        total_dias: int = 44,
        meta: int = None,
        andamento_diario=None,
        num_registros_reais: int = 0,
        leads_boletos_pendentes: int = 0,
        leads_em_negociacao: int = 0,
    ) -> dict:
        """
        Gera a projeção completa para um curso (inédito ou existente), usando o Ensemble ML
        calibrado pelo andamento do vestibular quando disponível.
        """
        periodo_norm, sigla_curso, turno_extenso = self.detectar_turno_e_curso(
            nome_curso=nome_curso, periodo_override=periodo
        )

        total_dias = max(1, int(total_dias))
        dia_atual = max(0, min(int(dia_atual), total_dias))
        inscritos_hoje = max(0.0, float(inscritos_hoje))

        ens = self._projetar_ensemble_andamento(
            periodo=periodo_norm,
            sigla_curso=sigla_curso,
            inscritos_hoje=inscritos_hoje,
            dia_atual=dia_atual,
            total_dias=total_dias,
            andamento_diario=andamento_diario,
            num_registros_reais=num_registros_reais,
        )

        proj_esperada = ens["projecao_esperada"]
        proj_pessimista = ens["cenario_pessimista"]
        proj_otimista = ens["cenario_otimista"]

        # Potencial adicional vindo do funil CRM do Odin (boletos emitidos e negociações ativas)
        potencial_leads = int(round(0.55 * leads_boletos_pendentes + 0.20 * leads_em_negociacao))
        proj_com_leads = int(round(proj_esperada)) + potencial_leads

        dias_restantes = max(1, total_dias - dia_atual)
        meta_efetiva = int(meta) if meta is not None else int(round(proj_esperada))
        deficit = max(0.0, meta_efetiva - inscritos_hoje)
        taxa_diaria = deficit / dias_restantes if meta is not None else 0.0

        # Taxa diária observada nos últimos 7 dias do andamento
        if andamento_diario and len(andamento_diario) >= 2:
            janela = min(7, len(andamento_diario) - 1)
            taxa_atual_7d = max(0.0, (andamento_diario[-1] - andamento_diario[-1 - janela]) / janela)
        else:
            taxa_atual_7d = inscritos_hoje / max(1, dia_atual) if dia_atual > 0 else 0.0

        if meta is not None:
            if proj_esperada >= meta:
                status = "Atingível (Dentro da Tendência)"
            elif proj_otimista >= meta:
                status = "Atenção (Exige Manter Ritmo Alto)"
            else:
                status = "Exige Aceleração Extra"
        else:
            status = "-"

        return {
            "curso": nome_curso,
            "sigla_identificada": sigla_curso,
            "periodo": periodo_norm,
            "turno_extenso": turno_extenso,
            "dia_atual": dia_atual,
            "total_dias": total_dias,
            "dias_restantes": dias_restantes,
            "inscritos_hoje": int(round(inscritos_hoje)),
            "ritmo_medio": ens["ritmo_medio"],
            "ritmo_medio_pct": round(ens["ritmo_medio"] * 100.0, 1),
            "ritmo_ajustado_ml": ens["ritmo_ajustado_ml"],
            "ritmo_ajustado_ml_pct": round(ens["ritmo_ajustado_ml"] * 100.0, 1),
            "cenario_pessimista": int(round(proj_pessimista)),
            "projecao_esperada": int(round(proj_esperada)),
            "projecao_pontual": int(round(ens["projecao_pontual"])),
            "projecao_kernel": int(round(ens["projecao_kernel"])),
            "projecao_ridge": int(round(ens["projecao_ridge"])),
            "cenario_otimista": int(round(proj_otimista)),
            "potencial_leads_quentes": potencial_leads,
            "projecao_com_leads": proj_com_leads,
            "meta": meta_efetiva,
            "deficit_meta": int(round(deficit)),
            "taxa_diaria_necessaria": round(taxa_diaria, 2),
            "taxa_diaria_atual_7d": round(taxa_atual_7d, 2),
            "status": status,
            "usou_andamento_ml": ens["usou_andamento_ml"],
            "curva_mais_similar": ens["curva_mais_similar"],
            "similaridade_max_pct": ens["similaridade_max_pct"],
            "top_curvas_similares": ens["top_curvas_similares"],
            "aceleracao_ratio": ens["aceleracao_ratio"],
            "aceleracao_label": ens["aceleracao_label"],
            "confianca_modelo": ens["confianca_modelo"],
            "_curva_aprendida_grid": ens["curva_aprendida_grid"],
        }

    def extrair_andamento_campanha(self, campaign):
        """
        Extrai a série temporal diária acumulada ('andamento') e indicadores de funil
        a partir de uma instância de VestibularCampaign do Odin.
        """
        dias = campaign.total_days
        elapsed = campaign.elapsed_days
        start = campaign.effective_start_date

        records = list(campaign.daily_records.all().order_by("date"))
        records_by_date = {}
        paid_before_start = 0
        sponsored_before_start = 0

        for r in records:
            if r.date < start:
                paid_before_start += r.new_paid
                sponsored_before_start += r.sponsored_released
            else:
                records_by_date[r.date] = r

        andamento_pagos = []
        andamento_patrocinados = []
        acum_p = paid_before_start
        acum_s = sponsored_before_start

        dias_observados = max(0, min(elapsed, dias))
        for i in range(dias_observados):
            d_curr = start + timedelta(days=i)
            if d_curr in records_by_date:
                acum_p += records_by_date[d_curr].new_paid
                acum_s += records_by_date[d_curr].sponsored_released
            andamento_pagos.append(acum_p)
            andamento_patrocinados.append(acum_s)

        # Caso existam registros em datas posteriores a elapsed (ex: testes futuros)
        total_paid = campaign.accumulated_paid
        if andamento_pagos and andamento_pagos[-1] < total_paid:
            andamento_pagos[-1] = total_paid

        leads_qs = campaign.leads.all()
        boletos_pendentes = leads_qs.filter(status="inscrito_nao_pago").count()
        em_negociacao = leads_qs.filter(status__in=["em_contato", "isencao"]).count()

        return {
            "dias_totais": dias,
            "dia_atual": dias_observados,
            "inscritos_pagos": total_paid,
            "patrocinados": campaign.accumulated_sponsored,
            "andamento_pagos": andamento_pagos,
            "andamento_patrocinados": andamento_patrocinados,
            "num_registros": len(records),
            "boletos_pendentes": boletos_pendentes,
            "em_negociacao": em_negociacao,
        }

    def projetar_campanha_odin(self, campaign, periodo_override: str = None) -> dict:
        """
        Integra o motor de Aprendizado de Máquina diretamente ao objeto VestibularCampaign do Odin,
        calculando tanto as projeções preditivas (Ensemble ML + Andamento) quanto a recomendação
        ótima de patrocínios (Yield Management).
        """
        dados = self.extrair_andamento_campanha(campaign)
        periodo, sigla, turno_extenso = self.detectar_turno_e_curso(
            nome_curso=campaign.course.name,
            nome_campanha=campaign.name,
            periodo_override=periodo_override,
        )

        if dados["dia_atual"] == 0:
            return {
                "progresso_esperado_pct": 0.0,
                "projecao_final_pagantes": 0,
                "projecao_pontual": 0,
                "projecao_kernel": 0,
                "projecao_ridge": 0,
                "cenario_pessimista": 0,
                "cenario_otimista": 0,
                "s_alvo": 0,
                "liberar_hoje": 0,
                "ritmo_medio_pct": 0.0,
                "ritmo_ajustado_ml_pct": 0.0,
                "meta": campaign.min_inscriptions,
                "deficit_meta": campaign.min_inscriptions,
                "taxa_diaria_necessaria": round(campaign.min_inscriptions / max(1, dados["dias_totais"]), 2),
                "taxa_diaria_atual_7d": 0.0,
                "status": "Campanha Não Iniciada",
                "usou_andamento_ml": False,
                "curva_mais_similar": "-",
                "similaridade_max_pct": 0.0,
                "top_curvas_similares": [],
                "aceleracao_ratio": 1.0,
                "aceleracao_label": "Aguardando Início",
                "confianca_modelo": 0.0,
                "periodo": periodo,
                "turno_extenso": turno_extenso,
                "sigla_identificada": sigla,
                "potencial_leads_quentes": dados["boletos_pendentes"],
                "projecao_com_leads": 0,
                "num_registros": dados["num_registros"],
            }

        res_ml = self.projetar(
            nome_curso=campaign.course.name,
            periodo=periodo,
            inscritos_hoje=dados["inscritos_pagos"],
            dia_atual=dados["dia_atual"],
            total_dias=dados["dias_totais"],
            meta=campaign.min_inscriptions,
            andamento_diario=dados["andamento_pagos"],
            num_registros_reais=dados["num_registros"],
            leads_boletos_pendentes=dados["boletos_pendentes"],
            leads_em_negociacao=dados["em_negociacao"],
        )

        projecao_final = res_ml["projecao_esperada"] if dados["inscritos_pagos"] > 0 else 0
        patrocinios_atuais = dados["patrocinados"]

        # Controlador de Yield Management (Cota Ótima de Patrocínios baseada na Projeção ML)
        yield_paid = campaign.yield_paid if campaign.yield_paid > 0 else 0.60
        yield_sponsored = campaign.yield_sponsored if campaign.yield_sponsored > 0 else 0.72
        # Se yield_paid/yield_sponsored foram cadastrados como taxas (>1 trata como valor monetário e usa default 0.60/0.72)
        taxa_presenca_pagos = yield_paid if yield_paid <= 1.0 else 0.60
        taxa_presenca_patroc = yield_sponsored if yield_sponsored <= 1.0 else 0.72

        gap_vagas = campaign.vacancies - (taxa_presenca_pagos * projecao_final)
        s_vagas = max(0.0, gap_vagas / taxa_presenca_patroc)
        s_volume = max(0.0, campaign.min_inscriptions - projecao_final)

        s_alvo = math.ceil(max(s_vagas, s_volume))
        liberar_hoje = max(0, s_alvo - patrocinios_atuais)

        return {
            "progresso_esperado_pct": res_ml["ritmo_ajustado_ml_pct"],
            "projecao_final_pagantes": int(projecao_final),
            "projecao_pontual": res_ml["projecao_pontual"],
            "projecao_kernel": res_ml["projecao_kernel"],
            "projecao_ridge": res_ml["projecao_ridge"],
            "cenario_pessimista": res_ml["cenario_pessimista"],
            "cenario_otimista": res_ml["cenario_otimista"],
            "s_alvo": s_alvo,
            "liberar_hoje": liberar_hoje,
            "ritmo_medio_pct": res_ml["ritmo_medio_pct"],
            "ritmo_ajustado_ml_pct": res_ml["ritmo_ajustado_ml_pct"],
            "meta": res_ml["meta"],
            "deficit_meta": res_ml["deficit_meta"],
            "taxa_diaria_necessaria": res_ml["taxa_diaria_necessaria"],
            "taxa_diaria_atual_7d": res_ml["taxa_diaria_atual_7d"],
            "status": res_ml["status"],
            "usou_andamento_ml": res_ml["usou_andamento_ml"],
            "curva_mais_similar": res_ml["curva_mais_similar"],
            "similaridade_max_pct": res_ml["similaridade_max_pct"],
            "top_curvas_similares": res_ml["top_curvas_similares"],
            "aceleracao_ratio": res_ml["aceleracao_ratio"],
            "aceleracao_label": res_ml["aceleracao_label"],
            "confianca_modelo": res_ml["confianca_modelo"],
            "periodo": periodo,
            "turno_extenso": turno_extenso,
            "sigla_identificada": sigla,
            "potencial_leads_quentes": res_ml["potencial_leads_quentes"],
            "projecao_com_leads": res_ml["projecao_com_leads"],
            "num_registros": dados["num_registros"],
            "_curva_aprendida_grid": res_ml["_curva_aprendida_grid"],
        }

    def gerar_dados_grafico_campanha(self, campaign, periodo_override: str = None) -> str:
        """
        Gera o JSON para o gráfico Chart.js do Odin:
          - labels: datas do período de inscrição
          - real_pagantes: evolução diária real dos inscritos pagos (até elapsed_days)
          - real_patrocinados: evolução diária real dos patrocínios (até elapsed_days)
          - projetada: curva esperada calibrada pelo modelo de ML (e ancorada no valor real até o fim)
          - cenario_otimista: cone superior de projeção ML do dia atual até o fim
          - cenario_pessimista: cone inferior de projeção ML do dia atual até o fim
        """
        try:
            dados = self.extrair_andamento_campanha(campaign)
            proj = self.projetar_campanha_odin(campaign, periodo_override=periodo_override)

            dias = dados["dias_totais"]
            elapsed = dados["dia_atual"]
            start = campaign.effective_start_date
            labels = [(start + timedelta(days=i)).strftime("%d/%m") for i in range(dias)]

            # Interpola a curva aprendida para os N dias da campanha
            curva_grid = proj.get("_curva_aprendida_grid")
            if curva_grid is None:
                curva_grid = np.linspace(0.0, 1.0, self.GRID_SIZE)
            t_dias_norm = np.linspace(0.0, 1.0, dias)
            f_curva = interp1d(self.grid, curva_grid, kind="linear", bounds_error=False, fill_value=(0.0, 1.0))
            curva_diaria = np.clip(f_curva(t_dias_norm), 0.0, 1.0)
            if curva_diaria[-1] > 0:
                curva_diaria = curva_diaria / curva_diaria[-1]

            alvo_esperado = proj["projecao_final_pagantes"]
            alvo_otimista = proj["cenario_otimista"]
            alvo_pessimista = proj["cenario_pessimista"]

            if alvo_esperado == 0:
                alvo_esperado = campaign.min_inscriptions or int(campaign.vacancies / 0.60)
                alvo_otimista = int(round(alvo_esperado * 1.2))
                alvo_pessimista = int(round(alvo_esperado * 0.8))

            linha_real_pagantes = []
            linha_real_patrocinados = []
            for i in range(dias):
                if i < len(dados["andamento_pagos"]):
                    linha_real_pagantes.append(int(dados["andamento_pagos"][i]))
                    linha_real_patrocinados.append(int(dados["andamento_patrocinados"][i]))
                else:
                    linha_real_pagantes.append(None)
                    linha_real_patrocinados.append(None)

            # Curva de referência ML (do dia 1 ao fim) + Cone de projeção ancorado no último dia observado
            linha_projetada = []
            linha_otimista = []
            linha_pessimista = []

            idx_ancora = max(0, elapsed - 1)
            val_ancora = dados["andamento_pagos"][idx_ancora] if dados["andamento_pagos"] else 0
            phi_ancora = float(curva_diaria[idx_ancora]) if dias > 0 else 0.0

            for i in range(dias):
                phi_i = float(curva_diaria[i])
                if elapsed == 0 or dados["inscritos_pagos"] == 0:
                    linha_projetada.append(int(round(phi_i * alvo_esperado)))
                    linha_otimista.append(int(round(phi_i * alvo_otimista)))
                    linha_pessimista.append(int(round(phi_i * alvo_pessimista)))
                elif i < idx_ancora:
                    # Antes do dia atual: mostra a trajetória esperada de referência para comparação visual
                    linha_projetada.append(int(round(phi_i * alvo_esperado)))
                    linha_otimista.append(None)
                    linha_pessimista.append(None)
                elif i == idx_ancora:
                    # Ponto de ancoragem (dia atual): conecta exatamente na medição real
                    linha_projetada.append(int(val_ancora))
                    linha_otimista.append(int(val_ancora))
                    linha_pessimista.append(int(val_ancora))
                else:
                    # Dias futuros (projeção do andamento até o encerramento)
                    fator_restante = (phi_i - phi_ancora) / max(1.0 - phi_ancora, 1e-4)
                    fator_restante = max(0.0, min(1.0, fator_restante))
                    v_esp = val_ancora + (alvo_esperado - val_ancora) * fator_restante
                    v_otim = val_ancora + (alvo_otimista - val_ancora) * fator_restante
                    v_pess = val_ancora + (alvo_pessimista - val_ancora) * fator_restante
                    linha_projetada.append(int(round(v_esp)))
                    linha_otimista.append(int(round(v_otim)))
                    linha_pessimista.append(int(round(v_pess)))

            return json.dumps({
                "labels": labels,
                "projetada": linha_projetada,
                "cenario_otimista": linha_otimista,
                "cenario_pessimista": linha_pessimista,
                "real_pagantes": linha_real_pagantes,
                "real_patrocinados": linha_real_patrocinados,
                "meta_inscriptions": campaign.min_inscriptions,
            })
        except Exception as e:
            return json.dumps({
                "labels": ["Erro de Cálculo", str(e)],
                "projetada": [0, 0],
                "cenario_otimista": [0, 0],
                "cenario_pessimista": [0, 0],
                "real_pagantes": [0, 0],
                "real_patrocinados": [0, 0],
            })

    def _construir_workbook_excel(self, resultados: list, dia_atual: int, total_dias: int):
        """Monta o Workbook openpyxl formatado com gráfico comparativo."""
        import openpyxl
        from openpyxl.chart import BarChart, Reference
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Projeções ML Vestibular"
        ws.views.sheetView[0].showGridLines = True

        font_title = Font(name="Segoe UI", size=15, bold=True, color="1E3A8A")
        font_sub = Font(name="Segoe UI", size=10, italic=True, color="475569")
        font_hdr = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        font_bold = Font(name="Segoe UI", size=11, bold=True, color="1E293B")

        fill_navy = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        fill_blue_soft = PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid")
        fill_green_soft = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
        fill_amber_soft = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")

        border_thin = Side(border_style="thin", color="CBD5E1")
        box = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)

        ws["B2"] = "Painel de Projeções: Aprendizado de Máquina & Andamento do Vestibular (Odin)"
        ws["B2"].font = font_title
        ws["B3"] = (
            f"Posição Referência: Dia {dia_atual} de {total_dias} | "
            "Calibrado por Ensemble ML (Trajetória Funcional + Ridge + Curvas Históricas por Turno)"
        )
        ws["B3"].font = font_sub

        headers = [
            "Curso / Campanha",
            "Turno",
            f"Inscritos (D-{dia_atual})",
            "Ritmo Histórico",
            "Cenário Pessimista",
            "Projeção ML (Esperada)",
            "Cenário Otimista",
            "Meta Definida",
            "Faltam p/ Meta",
            "Média Diária Necessária",
            "Diagnóstico Operacional",
            "Curva Histórica Mais Similar",
        ]

        for c_idx, h in enumerate(headers, start=2):
            cell = ws.cell(row=5, column=c_idx, value=h)
            cell.font = font_hdr
            cell.fill = fill_navy
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = box
        ws.row_dimensions[5].height = 28

        for r_i, r in enumerate(resultados, start=6):
            ws.cell(row=r_i, column=2, value=r["curso"]).alignment = Alignment(horizontal="left")
            ws.cell(row=r_i, column=3, value=r["turno_extenso"]).alignment = Alignment(horizontal="center")
            ws.cell(row=r_i, column=4, value=r["inscritos_hoje"]).alignment = Alignment(horizontal="right")

            c_ritmo = ws.cell(row=r_i, column=5, value=r["ritmo_medio"])
            c_ritmo.number_format = "0.0%"
            c_ritmo.alignment = Alignment(horizontal="right")

            ws.cell(row=r_i, column=6, value=r["cenario_pessimista"]).alignment = Alignment(horizontal="right")

            c_esp = ws.cell(row=r_i, column=7, value=r["projecao_esperada"])
            c_esp.alignment = Alignment(horizontal="right")
            c_esp.fill = fill_blue_soft
            c_esp.font = font_bold

            ws.cell(row=r_i, column=8, value=r["cenario_otimista"]).alignment = Alignment(horizontal="right")
            ws.cell(row=r_i, column=9, value=r["meta"]).alignment = Alignment(horizontal="right")
            ws.cell(row=r_i, column=10, value=r["deficit_meta"]).alignment = Alignment(horizontal="right")

            c_taxa = ws.cell(row=r_i, column=11, value=r["taxa_diaria_necessaria"])
            c_taxa.number_format = "0.00"
            c_taxa.alignment = Alignment(horizontal="right")

            c_st = ws.cell(row=r_i, column=12, value=r["status"])
            c_st.alignment = Alignment(horizontal="center")
            c_st.font = font_bold
            c_st.fill = fill_green_soft if "Atingível" in r["status"] else fill_amber_soft

            c_sim = ws.cell(row=r_i, column=13, value=r.get("curva_mais_similar", "-"))
            c_sim.alignment = Alignment(horizontal="center")

            for c_col in range(2, 14):
                c = ws.cell(row=r_i, column=c_col)
                c.border = box
                if c_col not in [7, 12] and r_i % 2 == 1:
                    c.fill = fill_zebra
            ws.row_dimensions[r_i].height = 22

        if resultados:
            tot_row = len(resultados) + 6
            ws.cell(row=tot_row, column=2, value="TOTAL GERAL").alignment = Alignment(horizontal="left")
            ws.cell(row=tot_row, column=3, value="-").alignment = Alignment(horizontal="center")
            ws.cell(row=tot_row, column=4, value=f"=SUM(D6:D{tot_row-1})").alignment = Alignment(horizontal="right")
            ws.cell(row=tot_row, column=5, value=f"=AVERAGE(E6:E{tot_row-1})").alignment = Alignment(horizontal="right")
            ws.cell(row=tot_row, column=5).number_format = "0.0%"
            ws.cell(row=tot_row, column=6, value=f"=SUM(F6:F{tot_row-1})").alignment = Alignment(horizontal="right")
            ws.cell(row=tot_row, column=7, value=f"=SUM(G6:G{tot_row-1})").alignment = Alignment(horizontal="right")
            ws.cell(row=tot_row, column=8, value=f"=SUM(H6:H{tot_row-1})").alignment = Alignment(horizontal="right")
            ws.cell(row=tot_row, column=9, value=f"=SUM(I6:I{tot_row-1})").alignment = Alignment(horizontal="right")
            ws.cell(row=tot_row, column=10, value=f"=SUM(J6:J{tot_row-1})").alignment = Alignment(horizontal="right")
            ws.cell(row=tot_row, column=11, value=f"=SUM(K6:K{tot_row-1})").alignment = Alignment(horizontal="right")
            ws.cell(row=tot_row, column=11).number_format = "0.00"
            ws.cell(row=tot_row, column=12, value="Monitoramento ML").alignment = Alignment(horizontal="center")
            ws.cell(row=tot_row, column=13, value="-").alignment = Alignment(horizontal="center")

            for c_col in range(2, 14):
                c = ws.cell(row=tot_row, column=c_col)
                c.border = box
                c.font = font_bold
                c.fill = fill_blue_soft
            ws.row_dimensions[tot_row].height = 24

            chart = BarChart()
            chart.type = "col"
            chart.style = 10
            chart.title = f"Comparativo: Atual (Dia {dia_atual}) vs. Projeção ML vs. Meta"
            chart.y_axis.title = "Inscrições Pagas"
            chart.x_axis.title = "Cursos"

            cats = Reference(ws, min_col=2, min_row=6, max_row=tot_row - 1)
            data_atual = Reference(ws, min_col=4, min_row=5, max_row=tot_row - 1)
            data_esp = Reference(ws, min_col=7, min_row=5, max_row=tot_row - 1)
            data_meta = Reference(ws, min_col=9, min_row=5, max_row=tot_row - 1)

            chart.add_data(data_atual, titles_from_data=True)
            chart.add_data(data_esp, titles_from_data=True)
            chart.add_data(data_meta, titles_from_data=True)
            chart.set_categories(cats)

            chart.width = 19
            chart.height = 11
            ws.add_chart(chart, f"B{tot_row + 2}")

        ws.column_dimensions["A"].width = 3
        ws.column_dimensions["B"].width = 34
        ws.column_dimensions["C"].width = 14
        ws.column_dimensions["D"].width = 18
        ws.column_dimensions["E"].width = 16
        ws.column_dimensions["F"].width = 18
        ws.column_dimensions["G"].width = 20
        ws.column_dimensions["H"].width = 18
        ws.column_dimensions["I"].width = 15
        ws.column_dimensions["J"].width = 15
        ws.column_dimensions["K"].width = 22
        ws.column_dimensions["L"].width = 28
        ws.column_dimensions["M"].width = 28

        return wb

    def gerar_relatorio_excel(
        self,
        lista_cursos_atuais: list,
        dia_atual: int,
        total_dias: int,
        caminho_saida="projecao_novos_cursos.xlsx",
    ):
        """
        Recebe a lista de cursos, processa as projeções ML (com suporte a 'andamento' opcional)
        e salva o Excel formatado com gráfico embutido.
        """
        import pandas as pd

        resultados = []
        for c in lista_cursos_atuais:
            res = self.projetar(
                nome_curso=c["nome"],
                periodo=c.get("periodo", None),
                inscritos_hoje=c["inscritos"],
                dia_atual=c.get("dia_atual", dia_atual),
                total_dias=c.get("total_dias", total_dias),
                meta=c.get("meta", None),
                andamento_diario=c.get("andamento", None),
                num_registros_reais=c.get("num_registros", 0),
            )
            resultados.append(res)

        wb = self._construir_workbook_excel(resultados, dia_atual=dia_atual, total_dias=total_dias)
        wb.save(caminho_saida)
        return pd.DataFrame([{k: v for k, v in r.items() if not k.startswith("_")} for r in resultados])

    def gerar_bytes_excel(self, resultados: list, dia_atual: int, total_dias: int) -> bytes:
        """Retorna os bytes do arquivo .xlsx em memória para download direto via HttpResponse no Django."""
        wb = self._construir_workbook_excel(resultados, dia_atual=dia_atual, total_dias=total_dias)
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    def obter_resumo_historico_csv(self) -> dict:
        """
        Retorna estatísticas e curvas históricas da base vestibular.csv para visualização
        na interface do Odin (comparação entre edições e curvas médias por turno).
        """
        curvas_m = [s["curve"] for s in self.series_treino if s["periodo"] == "M"]
        pesos_m = [s["peso"] for s in self.series_treino if s["periodo"] == "M"]
        curvas_n = [s["curve"] for s in self.series_treino if s["periodo"] == "N"]
        pesos_n = [s["peso"] for s in self.series_treino if s["periodo"] == "N"]

        media_m = np.average(curvas_m, axis=0, weights=pesos_m) if curvas_m else np.linspace(0, 1, self.GRID_SIZE)
        media_n = np.average(curvas_n, axis=0, weights=pesos_n) if curvas_n else np.linspace(0, 1, self.GRID_SIZE)

        # Amostra em 21 pontos (0%, 5%, 10%, ..., 100%) para o gráfico de curvas de aprendizado
        idx_sample = np.linspace(0, self.GRID_SIZE - 1, 21, dtype=int)
        labels_pct = [f"{int(round(self.grid[i] * 100))}%" for i in idx_sample]

        edicoes_resumo = []
        for s in self.series_treino:
            edicoes_resumo.append({
                "label": s["label"],
                "curso": s["curso"],
                "periodo": s["periodo"],
                "edicao": s["edicao"],
                "dias_totais": s["dias_totais"],
                "peso": s["peso"],
                "inscritos_finais": int(round(s["final"])),
                "p25": round(float(s["curve"][25]) * 100, 1),
                "p50": round(float(s["curve"][50]) * 100, 1),
                "p75": round(float(s["curve"][75]) * 100, 1),
                "p90": round(float(s["curve"][90]) * 100, 1),
            })

        return {
            "total_series_completas": len(self.series_treino),
            "total_series_parciais": len(self.series_parciais),
            "edicoes": edicoes_resumo,
            "chart_curvas_json": json.dumps({
                "labels": labels_pct,
                "curva_matutino": [round(float(media_m[i]) * 100, 1) for i in idx_sample],
                "curva_noturno": [round(float(media_n[i]) * 100, 1) for i in idx_sample],
            }),
        }


# Alias mantendo compatibilidade com o nome da classe em vestibularEstimator.py
PreditorVestibularGenerico = PreditorVestibularOdin
